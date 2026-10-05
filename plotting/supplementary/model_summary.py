"""Build S3-S6 from archived manuscript sources; no model reruns."""
from __future__ import annotations
import json
import re
import sys
from collections import Counter
import numpy as np
import pandas as pd
from scipy.stats import pearsonr
from scipy.io import loadmat
from matplotlib.colors import ListedColormap
from matplotlib.patches import Patch
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from common import *

from xlsx_helpers import strict_xlsx_sheets as _strict_xlsx_sheets

figures=[]
audit=[]
config=ROOT/'tables'/'supplementary'
gpath=ROOT/'tables/graphing/GEMStats/GEMStats.xlsx'
g=pd.read_excel(gpath,sheet_name='Sheet2').set_index('Models').loc[ORDER]
m=pd.read_excel(gpath,sheet_name='Sheet4').set_index('Models').loc[ORDER]
source_table(g.reset_index(),'S03_model_properties.tsv')
source_table(m.reset_index(),'S03_draft_curated_metrics.tsv')

def panel(ax, letter, title):
    ax.set_title(letter+'  '+title,loc='left',fontweight='bold',pad=10)

def strain_ticks(ax,axis='y'):
    if axis=='y':
        ax.set_yticks(range(21));ax.set_yticklabels(ORDER)
    else:
        ax.set_xticks(range(21));ax.set_xticklabels(ORDER,rotation=90)

# S3: genome associations and paired curation metrics.
fig,axs=plt.subplots(2,2,figsize=(6.5,5.6))
correlations=[]
for ax, y, title, ylabel in [(axs[0,0],g.Rxns,'Genome and reaction content','Curated reactions'),(axs[0,1],m.MetabolicGenes,'Genome and metabolic genes','Metabolic genes')]:
    x=m['Total Length (Mb)']; r,p=pearsonr(x,y)
    ax.scatter(x,y,c=[COLORS[s] for s in ORDER],s=22,edgecolor='white',linewidth=.3,zorder=3)
    xx=np.linspace(x.min(),x.max(),100);coef=np.polyfit(x,y,1)
    ax.plot(xx,np.polyval(coef,xx),color='#606060',lw=.9)
    ax.set_xlabel('Genome size (Mb)');ax.set_ylabel(ylabel)
    ax.text(.03,.96,'r = %.2f; P = %.2g\nn = 21 strains'%(r,p),transform=ax.transAxes,va='top',fontsize=7)
    panel(ax,'A' if ax is axs[0,0] else 'B',title)
    correlations.append({'relationship':title,'pearson_r':r,'two_sided_p':p,'n_strains':21})
for ax,metric,letter in [(axs[1,0],'F1','C'),(axs[1,1],'MCC','D')]:
    for s in ORDER:
        ax.plot([0,1],[m.loc[s,metric+'Draft'],m.loc[s,metric+'Final']],color=COLORS[s],alpha=.7,lw=.7)
        ax.scatter([0,1],[m.loc[s,metric+'Draft'],m.loc[s,metric+'Final']],color=COLORS[s],s=14,zorder=3)
    ax.set_xticks([0,1]);ax.set_xticklabels(['Draft','Curated']);ax.set_xlim(-.25,1.25)
    ax.set_ylim(-.05 if metric=='F1' else -.4,1.06);ax.set_yticks([0,.25,.5,.75,1] if metric=='F1' else [-.4,0,.5,1]);ax.set_ylabel(metric)
    panel(ax,letter,metric+' score')
handles=[Patch(color=v,label=k) for k,v in GROUP_COLORS.items()]
fig.legend(handles=handles,loc='lower center',ncol=3,frameon=False,fontsize=6.7,bbox_to_anchor=(.5,.005))
fig.subplots_adjust(left=.11,right=.98,top=.94,bottom=.18,hspace=.56,wspace=.39)
p3=[save(fig,'S03a_genome_and_curation')]
source_table(pd.DataFrame(correlations),'S03_correlations.tsv')

spath=EXTERNAL_DATA_ROOT/'Code/ModelCuration/$BioLogCuration/11_6thConfusionMatrix(TFA)/substrates_by_category5TFAm9.xlsx'
records=[]
for s in ORDER:
    d=pd.read_excel(spath,sheet_name='ttmodel_'+s)
    assert len(d)==29 and (d[['TP','FP','FN','TN']].sum(axis=1)==1).all()
    for _,row in d.iterrows():
        records.append({'strain':s,'condition':row.Substrate,'category':max(['TP','FP','FN','TN'],key=lambda c:row[c]),'provenance':'substrate-level archived scoring'})
    records.append({'strain':s,'condition':'PTYG','category':'TP','provenance':'recorded PTYG-positive aggregate condition'})
d=pd.DataFrame(records)
cats=['TN','TP','FP','FN'];subs=sorted(set(d.condition)-{'PTYG'})+['PTYG']
matrix=d.pivot(index='condition',columns='strain',values='category').loc[subs,ORDER]
aggregate=pd.crosstab(d.strain,d.category).reindex(index=ORDER,columns=cats,fill_value=0)
confpath=ROOT/'tables/graphing/ExperimentalGrowthData/SubstarateUtilization&ConfusionMetrics.xlsx'
reference=pd.read_excel(confpath,sheet_name='ConfusionMetrics').set_index('Model').loc[ORDER]
assert np.array_equal(aggregate[cats].values,reference[cats].values)
# The primary experimental workbook explicitly aligns names (row 2) to
# identifiers (row 3), rather than relying on names inferred from KEGG IDs.
cwbook=load_workbook(confpath,read_only=True,data_only=True)
cwrows=list(cwbook['Sheet3'].values)
secondary_path=EXTERNAL_DATA_ROOT/'Code/ExperimentalData/BioLogData/BioLogFinal.xlsx'
secondary=load_workbook(secondary_path,read_only=True,data_only=True)
secondary_names={r[1]:r[0] for r in secondary['Sheet2'].iter_rows(values_only=True) if r[1]}
# Resolve a typographical error in both experimental name sheets using the
# independent metabolite name paired with C00503 in a preserved KEGG model.
metabolite_name_path=EXTERNAL_DATA_ROOT/'Code/ModelCuration/$BioLogCuration/2_CuratingModels/1_flk/1_FixingFalseNegatives/keggModel.mat'
metabolite_name_model=loadmat(metabolite_name_path,squeeze_me=True,struct_as_record=False)['keggmodel']
erythritol_index=list(metabolite_name_model.mets).index('C00503')
assert metabolite_name_model.metNames[erythritol_index]=='Erythritol'
display_names={
 'C00042':'Succinate','C00062':'L-Arginine','C00065':'L-Serine',
 'C00079':'Phenylalanine','C00093':'D,L-α-Glycerol phosphate',
 'C00103':'Glucose 1-phosphate','C00109':'α-Ketobutyric acid',
 'C00134':'Putrescine','C00140':'N-Acetyl-D-glucosamine',
 'C00152':'L-Asparagine','C00156':'4-Hydroxybenzoic acid','C00181':'D-Xylose',
 'C00185':'D-Cellobiose','C00188':'L-Threonine','C00243':'α-D-Lactose',
 'C00333':'D-Galacturonic acid','C00334':'γ-Aminobutyric acid','C00369':'Glycogen',
 'C00392':'D-Mannitol','C00490':'Itaconic acid','C00497':'D-Malic acid',
 'C00503':'Erythritol','C00805':'2-Hydroxybenzoic acid',
 'C00973':'Cyclodextrin','C03383':'D-Galactonic acid γ-lactone',
 'C03752':'D-Glucosaminic acid','C05332':'Phenylethylamine',
 'C11625':'Tween 80','D05566':'Tween 40','PTYG':'PTYG'}
crosswalk=[]
for j in range(4,len(cwrows[1])):
    name=cwrows[1][j];identifier=cwrows[2][j]
    if name=='PTYG':identifier='PTYG'
    assert identifier in display_names
    if identifier in secondary_names:
        assert secondary_names[identifier]==name
    crosswalk.append({'condition_id':identifier,'source_condition_id':cwrows[2][j],'source_name':name,
      'display_name':display_names[identifier],
      'primary_source':str(confpath.relative_to(ROOT)),'source_sheet':'Sheet3',
      'name_cell':f'{get_column_letter(j+1)}2','id_cell':f'{get_column_letter(j+1)}3',
      'secondary_source':str(secondary_path.relative_to(ROOT)) if identifier in secondary_names else '',
      'secondary_sheet':'Sheet2' if identifier in secondary_names else '',
      'additional_name_source':str(metabolite_name_path.relative_to(ROOT)) if identifier=='C00503' else '',
      'additional_name_locator':'keggmodel.mets == C00503; paired metNames entry Erythritol' if identifier=='C00503' else '',
      'name_note':'Typography normalized; source spelling I-Erythriol corrected using the local metabolite name record.' if identifier=='C00503' else 'Typography normalized without changing the experimental substrate identity.'})
crosswalk=pd.DataFrame(crosswalk).set_index('condition_id').loc[subs].reset_index()
assert set(crosswalk.condition_id)==set(subs) and len(crosswalk)==30
source_table(crosswalk,'S03_condition_name_crosswalk.tsv')
audit.append({'analysis':'S3 condition display names','status':'verified against local name-to-identifier records',
 'details':'29 substrate names plus PTYG match the aligned Sheet3 name and identifier rows; 28 substrate mappings independently agree with BioLogFinal Sheet2. The recorded I-Erythriol spelling was corrected to Erythritol using C00503/metNames in the preserved KEGG model. PTYG uses the recorded medium name rather than its C0000 placeholder.'})
source_table(d,'S03_substrate_confusion_records.tsv')
source_table(aggregate.reset_index(),'S03_reconstructed_aggregate_confusion.tsv')
phen=pd.read_excel(confpath,sheet_name='Sheet1',header=1).set_index('Models')
discrepancies=[]
for s in ORDER:
    observed=int(pd.to_numeric(phen.loc[s].iloc[3:],errors='coerce').sum())
    scored=int(reference.loc[s,'TP']+reference.loc[s,'FN'])
    if observed!=scored:discrepancies.append({'strain':s,'displayed_positive_count':observed,'scoring_positive_count':scored,'resolution':'unresolved; scoring records retained explicitly'})
source_table(pd.DataFrame(discrepancies),'S03_phenotype_reference_discrepancies.tsv')
fig,ax=plt.subplots(figsize=(6.5,6.6))
cmap=ListedColormap(['#eeeeee','#527a8b','#d97732','#99446a'])
ax.imshow(matrix.replace(dict(zip(cats,range(4)))).values,cmap=cmap,vmin=-.5,vmax=3.5,aspect='auto',interpolation='none')
strain_ticks(ax,'x');ax.set_yticks(range(len(subs)));ax.set_yticklabels([display_names[s] for s in subs])
ax.set_xticks(np.arange(-.5,21,1),minor=True);ax.set_yticks(np.arange(-.5,len(subs),1),minor=True)
ax.grid(which='minor',color='white',lw=.6);ax.tick_params(which='minor',length=0)
panel(ax,'E','Substrate-level prediction agreement')
ax.set_ylabel('Growth substrate or medium')
fig.legend(handles=[Patch(color=cmap(i),label=lab) for i,lab in enumerate(['True negative','True positive','False positive','False negative'])],loc='lower center',ncol=2,frameon=False,bbox_to_anchor=(.53,.015))
fig.subplots_adjust(left=.33,right=.98,top=.94,bottom=.15)
p3.append(save(fig,'S03b_substrate_prediction_errors'))
figures.append(dict(number=3,title='Genome content and prediction agreement',pages=p3,sources=[str(gpath.relative_to(ROOT)),str(spath.relative_to(ROOT)),str(confpath.relative_to(ROOT)),str(secondary_path.relative_to(ROOT)),str(metabolite_name_path.relative_to(ROOT))],checks={'strains':21,'conditions_per_strain':30,'all_aggregate_confusion_counts_match_archived_table':True,'false_positives':int((d.category=='FP').sum()),'false_negatives':int((d.category=='FN').sum())},caption='(A, B) Genome size versus curated reaction content and metabolic-gene content for 21 strains. Lines are ordinary least-squares fits; annotations report Pearson correlations and two-sided P values. (C, D) Paired draft and curated F1 scores and Matthews correlation coefficients (MCC) from the archived scoring tables; each line denotes one strain. Colors identify the taxonomic groups used in the manuscript. (E) Substrate-level true-positive, true-negative, false-positive and false-negative records for 29 archived conditions, plus the recorded PTYG-positive condition. This reconstruction reproduces the aggregate confusion counts for all 21 strains. Agreement is assessed against the archived scoring reference used during curation, rather than an independent test set. The separate displayed phenotype matrix has one additional positive entry for each of cxy and lur; these unresolved reference differences are retained in the accompanying source table. No uncertainty interval is implied by the fitted lines.'))
audit.append({'analysis':'Substrate validation reference','status':'included with explicit unresolved discrepancy','details':discrepancies})

# S4: archived MEMOTE, reported thermodynamic coverage, relaxation counts.
mr=[];memfiles=[]
for s in ORDER:
    p=EXTERNAL_DATA_ROOT/'Code/ModelCuration/MEMOTE'/(s+'Memote.html')
    if not p.exists():p=EXTERNAL_DATA_ROOT/'Code/ModelCuration/MEMOTE'/(s+'.html')
    j=json.JSONDecoder().raw_decode(p.read_text().split('window.data = ',1)[1])[0]
    row={'strain':s,'overall':j['score']['total_score']*100,'timestamp':j['meta']['timestamp'],'memote_version':j['meta']['packages']['memote']}
    row.update({x['section']:x['score']*100 for x in j['score']['sections']});mr.append(row);memfiles.append(str(p.relative_to(ROOT)))
md=pd.DataFrame(mr).set_index('strain')
source_table(md.reset_index(),'S04_historical_memote_scores.tsv')
rel=pd.read_csv(config/'thermo_relaxations.tsv',sep='\t')
source_table(rel,'S04_thermodynamic_relaxation_records.tsv')
source_table(g[['%ThermoMetabolites','%ThermoReactions']].reset_index(),'S04_reported_thermodynamic_coverage.tsv')
ptyg=rel[rel.medium=='PTYG'].set_index('strain_id').loc[ORDER]
minimal=rel[rel.medium=='M9'].set_index('strain_id')
assert len(ptyg)==21
cats=['overall','consistency','annotation_met','annotation_rxn','annotation_gene','annotation_sbo']
fig,axs=plt.subplots(1,3,figsize=(6.5,6.0),gridspec_kw={'width_ratios':[2.9,1.65,1.15]})
im=axs[0].imshow(md[cats].values,cmap='Blues',vmin=0,vmax=100,aspect='auto')
strain_ticks(axs[0]);axs[0].set_xticks(range(6));axs[0].set_xticklabels(['Overall','Consistency','Metabolite ann.','Reaction ann.','Gene ann.','SBO ann.'],rotation=65,ha='right')
for i in range(21):
    for k in range(6):axs[0].text(k,i,'%.1f'%md.iloc[i][cats[k]],ha='center',va='center',fontsize=6.1,color='white' if md.iloc[i][cats[k]]>65 else '#222222')
panel(axs[0],'A','MEMOTE scores (%)')
y=np.arange(21)
axs[1].scatter(g['%ThermoMetabolites'],y,s=17,color='#527a8b',label='Metabolites')
axs[1].scatter(g['%ThermoReactions'],y,s=17,marker='s',color='#d97732',label='Reactions')
axs[1].set_ylim(20.5,-.5);axs[1].set_xlim(0,100);axs[1].set_yticks([]);axs[1].set_xlabel('Reported coverage (%)');axs[1].set_xticks([0,50,100]);axs[1].legend(loc='upper center',bbox_to_anchor=(.5,-.17),frameon=False,fontsize=6.5)
panel(axs[1],'B','Thermodynamic\ncoverage')
axs[2].barh(y,ptyg.reaction_count,color='#777777',height=.65);axs[2].set_ylim(20.5,-.5);axs[2].set_yticks([]);axs[2].set_xlim(0,43);axs[2].set_xlabel('Reaction count')
for i,v in enumerate(ptyg.reaction_count):axs[2].text(v+.7,i,str(v),va='center',fontsize=6)
panel(axs[2],'C','Relaxed\nconstraints')
fig.subplots_adjust(left=.075,right=.985,bottom=.28,top=.93,wspace=.25)
p4=save(fig,'S04_quality_and_thermodynamics')
figures.append(dict(number=4,title='Model quality and thermodynamic coverage',pages=[p4],sources=memfiles+[str(gpath.relative_to(ROOT)),str((config/'thermo_relaxations.tsv').relative_to(ROOT))],checks={'historical_reports':21,'memote_mean_percent':float(md.overall.mean()),'maximum_relaxed_reactions':int(ptyg.reaction_count.max())},caption='(A) Overall and category scores extracted from the 21 archived MEMOTE 0.14.0 reports. Annotation categories concern metabolites, reactions, genes and Systems Biology Ontology (SBO) terms. All scores use a common 0–100% scale; these historical reports describe the manuscript model generation. (B) Reported percentages of metabolites and reactions with thermodynamic data from the model-properties workbook. (C) Numbers of reactions whose thermodynamic constraints were relaxed in the archived PTYG models; labels give counts, not the magnitudes of relaxation. The corresponding M9 records are provided separately for the 19 strains with retained M9 results. Reaction identities, source workspaces and the documented cohm zero-relaxation correction are supplied with the data. These diagnostics describe model annotation and thermodynamic parameterization; they do not constitute experimental validation.'))

# S5: retain recorded evidence, including blank scores.
gfpath=ROOT/'tables/supplementary/GFUnified.xlsx';sheets=_strict_xlsx_sheets(gfpath);gr=[]
for s in ORDER:
    rows=sheets[s];head=rows[0]
    for k,row in enumerate(rows[1:],2):
        if not row or not re.match(r'^(?:R\d|rat\d)',row[0]):continue
        entry={head[i]:row[i] if i<len(row) else '' for i in range(len(head))}
        if not entry['Rxns']:continue
        entry.update(strain=s,source_row=k)
        entry['numeric_score']=pd.to_numeric(entry.get('Score',''),errors='coerce')
        gr.append(entry)
gd=pd.DataFrame(gr);assert len(gd)==577 and gd.numeric_score.notna().sum()==545
source_table(gd,'S05_gapfill_evidence_records.tsv')
source_table(pd.DataFrame(sheets['Scoring'],columns=['recorded_evidence_rule','score_or_increment']),'S05_recorded_scoring_rubric.tsv')
gd['plot_score']=gd.numeric_score.fillna(-1).astype(int)
count=pd.crosstab(gd.strain,gd.plot_score).reindex(index=ORDER,columns=[2,3,4,5,6,-1],fill_value=0)
source_table(count.rename(columns={-1:'unscored'}).reset_index(),'S05_evidence_score_counts.tsv')
fig,axs=plt.subplots(1,2,figsize=(6.5,5.5),gridspec_kw={'width_ratios':[3.7,1.5]})
palette=['#eff3ff','#c6dbef','#9ecae1','#4292c6','#08519c','#bdbdbd'];left=np.zeros(21)
for cat,col in zip(count.columns,palette):
    vals=count[cat].values/count.sum(axis=1).values*100
    axs[0].barh(range(21),vals,left=left,color=col,height=.74,label='Unscored' if cat==-1 else str(cat));left+=vals
strain_ticks(axs[0]);axs[0].set_ylim(20.7,-.7);axs[0].set_xlim(0,100);axs[0].set_xlabel('Reactions assessed (%)');panel(axs[0],'A','Evidence scores')
axs[0].legend(loc='upper center',bbox_to_anchor=(.52,-.13),ncol=3,title='Score',frameon=False,fontsize=7,title_fontsize=7)
axs[1].barh(range(21),count.sum(axis=1),height=.74,color='#777777');axs[1].set_ylim(20.7,-.7);axs[1].set_yticks([]);axs[1].set_xlim(0,count.sum(axis=1).max()*1.18);axs[1].set_xlabel('Number of reactions');panel(axs[1],'B','Gap-filling reactions')
for i,v in enumerate(count.sum(axis=1)):axs[1].text(v+.6,i,str(v),va='center',fontsize=7)
fig.subplots_adjust(left=.09,right=.98,top=.93,bottom=.19,wspace=.32)
p5=save(fig,'S05_gapfill_evidence_scores')
figures.append(dict(number=5,title='Recorded evidence for gap-filled reactions',pages=[p5],sources=[str(gfpath.relative_to(ROOT))],checks={'reaction_entries':577,'scored_entries':545,'unscored_entries':32,'score_distribution':{str(k):int(v) for k,v in gd.numeric_score.value_counts().items()}},caption='(A) Distribution of recorded evidence scores within each strain in the gap-filling audit workbook. (B) Number of recorded reaction entries per strain. Across the collection, 545 of 577 entries have numeric scores and 32 are unscored; missing scores remain a separate category. The archived rubric assigns 5 for an enzyme classification (EC) reported in the species, 4 in a closely related species, 3 in the genus, 2 for a partial three-level EC match in the species or a closely related species, and 1 for a partial match in the genus; presence of the gene in the relaxed reconstruction adds one point. No score-1 entries occur in this audit. Scores are reported as recorded, without inferring evidence from blank or displaced annotation fields. Counts describe the evidence audit and should not be equated with final integrated gap-fill totals, which differ across source stages. Reaction-level annotations and the original scoring rubric accompany this figure.'))
audit.append({'analysis':'Gap-fill evidence counts','status':'source generations kept separate','details':'577 evidence entries; 545 scored, 32 unscored. Do not equate audit rows to final integrated gap-fill totals; lgu has conflicting stage-specific totals.'})

# S6: audit exact retained row selection used by archived summary.
ipath=ROOT/'tables/graphing/iMM/IMMDataNames_with_classification_and_stats.xlsx';ir=[];comp=[]
for s in ORDER:
    df=pd.read_excel(ipath,sheet_name=s+'_summary')
    vals=df.drop(columns=['Row'],errors='ignore')
    mask=vals.applymap(lambda x:isinstance(x,str) and x.strip().lower()=='none').any(axis=1)
    if not mask.any():mask[:]=True
    mask &= ~vals.isna().all(axis=1)
    for j,row in vals.iterrows():
        compounds=[v.strip() for v in row if isinstance(v,str) and v.strip() and v.strip().lower()!='none']
        water=[v for v in compounds if v.lower()=='h2o']
        unique=list(dict.fromkeys(compounds));without=[v for v in unique if v.lower()!='h2o']
        rid=str(df.loc[j,'Row']) if 'Row' in df else str(j+2)
        ir.append({'strain':s,'source_row':j+2,'alternative_id':rid,'retained_in_archived_summary':bool(mask[j]),'compound_count_with_water':len(unique),'compound_count_without_water':len(without),'water_present':bool(water),'duplicate_compound_entries':len(compounds)-len(unique)})
        for v in unique:comp.append({'strain':s,'alternative_id':rid,'retained_in_archived_summary':bool(mask[j]),'compound':v,'is_water':v.lower()=='h2o'})
idd=pd.DataFrame(ir);ret=idd[idd.retained_in_archived_summary].copy();summary=[]
for s in ORDER:
    t=ret[ret.strain==s];minimum=t.compound_count_without_water.min()
    summary.append({'strain':s,'retained_alternatives':len(t),'min_cardinality_excluding_water':minimum,'max_cardinality_excluding_water':t.compound_count_without_water.max(),'median_cardinality_excluding_water':t.compound_count_without_water.median(),'at_retained_minimum_cardinality':int((t.compound_count_without_water==minimum).sum()),'water_containing_alternatives':int(t.water_present.sum())})
summary=pd.DataFrame(summary).set_index('strain');assert summary.loc['rfa','retained_alternatives']==40 and summary.loc['vpe','retained_alternatives']==120
source_table(idd,'S06_all_alternative_cardinalities.tsv');source_table(pd.DataFrame(comp),'S06_complete_nutrient_alternatives.tsv');source_table(summary.reset_index(),'S06_retained_alternative_summary.tsv')
fig,axs=plt.subplots(1,2,figsize=(6.5,5.5),gridspec_kw={'width_ratios':[1,1.1]})
axs[0].barh(range(21),summary.retained_alternatives,color='#527a8b',height=.68)
for i,s in enumerate(ORDER):
    n=summary.loc[s,'retained_alternatives'];axs[0].text(n+1.8,i,str(n),va='center',fontsize=7)
strain_ticks(axs[0]);axs[0].set_ylim(20.7,-.7);axs[0].set_xlim(0,142);axs[0].set_xlabel('Number of combinations');panel(axs[0],'A','Nutrient combinations')
for i,s in enumerate(ORDER):
    t=ret[ret.strain==s];freq=t.compound_count_without_water.value_counts()
    axs[1].plot([freq.index.min(),freq.index.max()],[i,i],color='#888888',lw=1,zorder=1)
    for v,n in freq.items():
        axs[1].scatter(v,i,s=24+70*n/len(t),color='#527a8b',edgecolor='white',lw=.5,zorder=3)
        dx=-8 if len(freq)>1 and v==freq.index.min() else 8
        axs[1].annotate(str(n),(v,i),xytext=(dx,0),textcoords='offset points',ha='right' if dx<0 else 'left',va='center',fontsize=6.4)
axs[1].set_ylim(20.7,-.7);axs[1].set_yticks([]);axs[1].set_xlim(6,18);axs[1].set_xticks([6,8,10,12,14,16,18]);axs[1].set_xlabel('Number of nutrients (water excluded)');panel(axs[1],'B','Nutrients per combination')
fig.subplots_adjust(left=.09,right=.98,top=.93,bottom=.12,wspace=.31)
p6=save(fig,'S06_nutrient_alternative_cardinality')
figures.append(dict(number=6,title='Nutrient combinations and their sizes',pages=[p6],sources=[str(ipath.relative_to(ROOT)),'RawData/iMM/plot_imm_core_variable_heatmap.py'],checks={'strains':21,'retained_alternatives_total':int(len(ret)),'rfa_alternatives':40,'vpe_alternatives':120,'duplicate_compound_entries':int(idd.duplicate_compound_entries.sum())},caption='(A) Numbers of retained nutrient alternatives used in the archived frequency summary. (B) Numbers of distinct supplied compounds per retained alternative after excluding water (H2O). Bubble area scales with the within-strain fraction of alternatives at a given cardinality; adjacent labels give the number of alternatives, and lines connect the observed minimum and maximum. Retention follows the archived workbook convention: rows containing a literal “none” padding entry are selected; if no such rows exist for a strain, all nonempty rows are retained. These rows are not uniformly strict minimum-cardinality solutions: lgu and lur include larger alternatives, and excluding water produces two cardinalities for cuvk, cxy and mgk and changes one lur entry. Accordingly, counts describe retained enumerated alternatives, not exhaustive solution numbers. The rfa count is recomputed from its 40 retained rows, correcting the stale zero in the workbook’s upper summary. Complete compositions, row identities, selection flags and cardinalities with and without water are provided as supplementary data.'))
audit.append({'analysis':'Nutrient alternative definition','status':'included with exact archived selection and cardinality audit','details':'Archive none-padding rule reproduced; strict minimality not assumed. rfa upper summary zero is stale; 40 rows recovered. Water excluded from plotted sizes, retained in complete compositions.'})

manifest('model_diagnostics',figures,audit)
print(json.dumps({'figures':[x['number'] for x in figures],'checks':[x['checks'] for x in figures],'phenotype_discrepancies':discrepancies},indent=2))
